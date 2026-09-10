import os
import openpyxl
import pandas as pd

folder = r"C:\Users\moonh\Desktop\Anythingllm 자료폴더\조류 통계"
f_patrol = os.path.join(folder, "일일점검(17_20년).xlsx")

print("--- Inspecting 일일점검(17_20년).xlsx ---")
if os.path.exists(f_patrol):
    wb = openpyxl.load_workbook(f_patrol, read_only=True)
    print("Sheet names:", wb.sheetnames)
    ws = wb[wb.sheetnames[0]]
    for r in range(1, 10):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, 20)]
        print(f"Row {r}:", row_vals)
    wb.close()
else:
    print("File not found!")
