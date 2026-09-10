import os
import openpyxl
import pandas as pd

folder = r"C:\Users\moonh\Desktop\Anythingllm 자료폴더\조류 통계"

# 1. Carcass 17-20
wb_c = openpyxl.load_workbook(os.path.join(folder, "사체수거(17_20년).xlsx"), data_only=True)
ws_c = wb_c[wb_c.sheetnames[0]]
c_rows = []
for r in range(4, ws_c.max_row + 1):
    d_val = ws_c.cell(row=r, column=5).value
    if not d_val: continue
    c_rows.append({
        'Date': str(d_val)[:10],
        'Year': str(d_val)[:4],
        'Time': str(ws_c.cell(row=r, column=6).value or '')[:8],
        'Location': str(ws_c.cell(row=r, column=7).value or '').strip(),
        'Species': str(ws_c.cell(row=r, column=8).value or '').strip(),
        'Weather': str(ws_c.cell(row=r, column=9).value or '').strip(),
        'Status': str(ws_c.cell(row=r, column=10).value or '').strip(),
        'Remark': str(ws_c.cell(row=r, column=11).value or '').strip()
    })
df_c = pd.DataFrame(c_rows)

# 2. Live Animals 17-20
wb_l = openpyxl.load_workbook(os.path.join(folder, "생동물(17_20년).xlsx"), data_only=True)
ws_l = wb_l[wb_l.sheetnames[0]]
l_rows = []
for r in range(4, ws_l.max_row + 1):
    d_val = ws_l.cell(row=r, column=3).value
    if not d_val: continue
    l_rows.append({
        'Date': str(d_val)[:10],
        'Year': str(d_val)[:4],
        'Time': str(d_val)[11:16],
        'Species': str(ws_l.cell(row=r, column=6).value or '').strip(),
        'Event': str(ws_l.cell(row=r, column=5).value or '').strip(),
        'Action': str(ws_l.cell(row=r, column=7).value or '').strip(),
        'Location': str(ws_l.cell(row=r, column=8).value or '').strip(),
        'Action_Detail': str(ws_l.cell(row=r, column=9).value or '').strip(),
        'Result': str(ws_l.cell(row=r, column=10).value or '').strip()
    })
df_l = pd.DataFrame(l_rows)

# 3. Bird Strike 17-20
wb_s = openpyxl.load_workbook(os.path.join(folder, "조류충돌(17_20년).xlsx"), data_only=True)
ws_s = wb_s['목록']
s_rows = []
for r in range(4, ws_s.max_row + 1):
    d_val = ws_s.cell(row=r, column=3).value
    if not d_val: continue
    s_rows.append({
        'Date': str(d_val)[:10],
        'Year': str(d_val)[:4],
        'Time': str(ws_s.cell(row=r, column=4).value or '')[:8],
        'Species': str(ws_s.cell(row=r, column=5).value or '').strip(),
        'Airline': str(ws_s.cell(row=r, column=6).value or '').strip(),
        'Flight': str(ws_s.cell(row=r, column=7).value or '').strip(),
        'Model': str(ws_s.cell(row=r, column=8).value or '').strip(),
        'Route': str(ws_s.cell(row=r, column=9).value or '').strip(),
        'Runway': str(ws_s.cell(row=r, column=10).value or '').strip(),
        'Altitude': ws_s.cell(row=r, column=11).value,
        'Phase': str(ws_s.cell(row=r, column=12).value or '').strip(),
        'Part': str(ws_s.cell(row=r, column=13).value or '').strip(),
        'Damage': str(ws_s.cell(row=r, column=14).value or 'N').strip(),
        'Position': str(ws_s.cell(row=r, column=15).value or '').strip(),
        'Found': str(ws_s.cell(row=r, column=16).value or '').strip(),
        'DNA': str(ws_s.cell(row=r, column=17).value or '').strip()
    })
df_s = pd.DataFrame(s_rows)

with open(r"C:\Users\moonh\Documents\iiac-BAT\scratch\stats_summary.txt", "w", encoding="utf-8") as out:
    out.write("=== 1. 생동물 (17_20년.xlsx) 총 101건 분석 ===\n")
    out.write(f"연도별 건수:\n{df_l['Year'].value_counts().sort_index().to_string()}\n\n")
    out.write(f"출몰 동물 종별 건수:\n{df_l['Species'].value_counts().to_string()}\n\n")
    out.write(f"조치 유형:\n{df_l['Action'].value_counts().head(10).to_string()}\n\n")
    out.write("주요 출몰 지역 TOP 10:\n" + df_l['Location'].value_counts().head(10).to_string() + "\n\n")
    
    out.write("\n=== 2. 조류충돌 (17_20년.xlsx) 총 325건 분석 ===\n")
    out.write(f"연도별 충돌 건수:\n{df_s['Year'].value_counts().sort_index().to_string()}\n\n")
    out.write(f"기체 손상(Damage Y) 연도별 현황:\n{df_s[df_s['Damage']=='Y']['Year'].value_counts().sort_index().to_string()}\n\n")
    out.write(f"손상률: {len(df_s[df_s['Damage']=='Y'])}/{len(df_s)} ({len(df_s[df_s['Damage']=='Y'])/len(df_s)*100:.1f}%)\n\n")
    out.write("활주로별 충돌 분포:\n" + df_s['Runway'].value_counts().to_string() + "\n\n")
    out.write("비행 단계(Flight Phase) 분포:\n" + df_s['Phase'].value_counts().to_string() + "\n\n")
    out.write("충돌 부위(Part) TOP 10:\n" + df_s['Part'].value_counts().head(10).to_string() + "\n\n")
    out.write("식별 조류종 TOP 10:\n" + df_s['Species'].value_counts().head(10).to_string() + "\n\n")
    
    out.write("\n=== 3. 사체수거 (17_20년.xlsx) 총 66건 분석 ===\n")
    out.write(f"연도별 수거 건수:\n{df_c['Year'].value_counts().sort_index().to_string()}\n\n")
    out.write("수거 조류종 분포:\n" + df_c['Species'].value_counts().head(10).to_string() + "\n\n")
    out.write("사체 상태 분포:\n" + df_c['Status'].value_counts().to_string() + "\n\n")
    out.write("수거 위치 TOP 10:\n" + df_c['Location'].value_counts().head(10).to_string() + "\n\n")

print("Analysis written to stats_summary.txt successfully!")
