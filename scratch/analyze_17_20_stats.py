import os
import openpyxl
import pandas as pd

folder = r"C:\Users\moonh\Desktop\Anythingllm 자료폴더\조류 통계"

# 1. Carcass 17-20
f_carcass = os.path.join(folder, "사체수거(17_20년).xlsx")
wb_c = openpyxl.load_workbook(f_carcass, data_only=True)
ws_c = wb_c[wb_c.sheetnames[0]]
c_rows = []
for r in range(4, ws_c.max_row + 1):
    d_val = ws_c.cell(row=r, column=5).value
    if not d_val: continue
    t_val = ws_c.cell(row=r, column=6).value
    loc = ws_c.cell(row=r, column=7).value
    sp = ws_c.cell(row=r, column=8).value
    weather = ws_c.cell(row=r, column=9).value
    status = ws_c.cell(row=r, column=10).value
    remark = ws_c.cell(row=r, column=11).value
    c_rows.append({
        'Date': str(d_val)[:10],
        'Time': str(t_val)[:8],
        'Location': loc,
        'Species': sp,
        'Weather': weather,
        'Status': status,
        'Remark': remark
    })
df_c = pd.DataFrame(c_rows)
print("=== Carcass (17_20년) ===")
print("Total rows:", len(df_c))
print("Year breakdown:")
print(df_c['Date'].str[:4].value_counts().sort_index())
print("Species top 10:")
print(df_c['Species'].value_counts().head(10))

# 2. Live Animals 17-20
f_live = os.path.join(folder, "생동물(17_20년).xlsx")
wb_l = openpyxl.load_workbook(f_live, data_only=True)
ws_l = wb_l[wb_l.sheetnames[0]]
l_rows = []
for r in range(4, ws_l.max_row + 1):
    d_val = ws_l.cell(row=r, column=3).value
    if not d_val: continue
    flight = ws_l.cell(row=r, column=4).value
    event = ws_l.cell(row=r, column=5).value
    sp = ws_l.cell(row=r, column=6).value
    action = ws_l.cell(row=r, column=7).value
    loc = ws_l.cell(row=r, column=8).value
    action_detail = ws_l.cell(row=r, column=9).value
    result = ws_l.cell(row=r, column=10).value
    l_rows.append({
        'DateTime': str(d_val)[:19],
        'Date': str(d_val)[:10],
        'Species': sp,
        'Event': event,
        'Action': action,
        'Location': loc,
        'Action_Detail': action_detail,
        'Result': result
    })
df_l = pd.DataFrame(l_rows)
print("\n=== Live Animals (17_20년) ===")
print("Total rows:", len(df_l))
print("Year breakdown:")
print(df_l['Date'].str[:4].value_counts().sort_index())
print("Species breakdown:")
print(df_l['Species'].value_counts())
print("Action breakdown:")
print(df_l['Action'].value_counts())

# 3. Bird Strike 17-20
f_strike = os.path.join(folder, "조류충돌(17_20년).xlsx")
wb_s = openpyxl.load_workbook(f_strike, data_only=True)
ws_s = wb_s['목록']
s_rows = []
for r in range(4, ws_s.max_row + 1):
    d_val = ws_s.cell(row=r, column=3).value
    if not d_val: continue
    t_val = ws_s.cell(row=r, column=4).value
    sp = ws_s.cell(row=r, column=5).value
    airline = ws_s.cell(row=r, column=6).value
    flight = ws_s.cell(row=r, column=7).value
    model = ws_s.cell(row=r, column=8).value
    route = ws_s.cell(row=r, column=9).value
    runway = ws_s.cell(row=r, column=10).value
    altitude = ws_s.cell(row=r, column=11).value
    phase = ws_s.cell(row=r, column=12).value
    part = ws_s.cell(row=r, column=13).value
    damage = ws_s.cell(row=r, column=14).value
    pos = ws_s.cell(row=r, column=15).value
    found = ws_s.cell(row=r, column=16).value
    dna = ws_s.cell(row=r, column=17).value
    s_rows.append({
        'Date': str(d_val)[:10],
        'Time': str(t_val)[:8],
        'Species': sp,
        'Airline': airline,
        'Flight': flight,
        'Model': model,
        'Route': route,
        'Runway': runway,
        'Altitude': altitude,
        'Phase': phase,
        'Part': part,
        'Damage': damage,
        'Position': pos,
        'Found': found,
        'DNA': dna
    })
df_s = pd.DataFrame(s_rows)
print("\n=== Bird Strike (17_20년) ===")
print("Total rows:", len(df_s))
print("Year breakdown:")
print(df_s['Date'].str[:4].value_counts().sort_index())
print("Damage breakdown:")
print(df_s['Damage'].value_counts())
print("Top 10 species:")
print(df_s['Species'].value_counts().head(10))
print("Runway breakdown:")
print(df_s['Runway'].value_counts())
