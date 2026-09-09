import os
import re
import sys
import io
import pandas as pd
import openpyxl
from datetime import datetime, time

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

FILE_PATROL_2021 = r"C:\Users\moonh\Desktop\Anythingllm 자료폴더\조류 통계\일일점검(21년).xlsx"
FILE_STRIKE_21_23 = r"C:\Users\moonh\Desktop\Anythingllm 자료폴더\조류 통계\조류충돌(21_23년).xlsx"
FILE_CARCASS_21_23 = r"C:\Users\moonh\Desktop\Anythingllm 자료폴더\조류 통계\사체수거(21_23년).xlsx"

OUT_PATROL_CSV = r"c:\Users\moonh\Documents\iiac-BAT\_data_raw\resampled\2021_total_resampled.csv"
OUT_STRIKE_CSV = r"c:\Users\moonh\Documents\iiac-BAT\_data_raw\resampled\2021_birdstrike_resampled.csv"
OUT_CARCASS_CSV = r"c:\Users\moonh\Documents\iiac-BAT\_data_raw\resampled\2021_carcass_resampled.csv"

def classify_zone(group_map, area_major, area_minor, raw_loc=""):
    s = f"{str(group_map)} {str(area_major)} {str(area_minor)} {str(raw_loc)}".lower()
    
    if any(k in s for k in ['as-34', '제4활주로 (남)', '34l', '34 남', 'ptw', 'p1', 'p2', 'p3', 'p4', 'p5', 'p6', 'p7', 'p8']):
        return 'ZONE-R4-S'
    elif any(k in s for k in ['as-16', '제4활주로 (북)', '16r', '16 북', 'p9', 'p10', 'p11', 'p12']):
        return 'ZONE-R4-N'
    elif any(k in s for k in ['as-34r', '제3활주로 (남)', '34r', '34r 남', 'w', 'u']):
        return 'ZONE-R3-S'
    elif any(k in s for k in ['as-16l', '제3활주로 (북)', '16l', '16l 북', 'n']):
        return 'ZONE-R3-N'
    elif any(k in s for k in ['as-33l', '제2활주로 (남)', '33l', '33l 남', 'd5', 'd4', 'd3']):
        return 'ZONE-R2-S'
    elif any(k in s for k in ['as-15r', '제2활주로 (북)', '15r', '15r 북', 'b']):
        return 'ZONE-R2-N'
    elif any(k in s for k in ['as-33r', '제1활주로 (남)', '33r', '33r 남', 'c', 'e', 'k']):
        return 'ZONE-R1-S'
    elif any(k in s for k in ['as-15l', '제1활주로 (북)', '15l', '15l 북', 'a']):
        return 'ZONE-R1-N'
    elif any(k in s for k in ['4활주로', '제4활주로']):
        return 'ZONE-R4-AIRSIDE'
    elif any(k in s for k in ['3활주로', '제3활주로']):
        return 'ZONE-R3-AIRSIDE'
    elif any(k in s for k in ['2활주로', '제2활주로']):
        return 'ZONE-R2-AIRSIDE'
    elif any(k in s for k in ['1활주로', '제1활주로']):
        return 'ZONE-R1-AIRSIDE'
    elif any(k in s for k in ['600', '612', '641', '648', '673', '841', '851', 'aact', '화물']):
        return 'ZONE-R2-APRON'
    elif any(k in s for k in ['200', '217', '251', '253', 't2']):
        return 'ZONE-T2-APRON'
    elif any(k in s for k in ['11번', '28번', '51번', 't1', '탑승동']):
        return 'ZONE-T1-APRON'
    elif any(k in s for k in ['배수로', '유수지', '암거']):
        return 'ZONE-DR-ISLAND'
    elif any(k in s for k in ['ls', '랜드사이드', '공사', 'ibc', '국제업무']):
        return 'ZONE-LANDSIDE'
    elif any(k in s for k in ['외곽', '철책', '초소', '울타리', '남측']):
        return 'ZONE-PERI-SOUTH'
    else:
        return 'ZONE-AIRSIDE-GENERAL'

def determine_credibility(species, activity, count, hour, dt_str):
    sp = str(species).strip()
    act = str(activity).strip()
    
    if any(k in act for k in ['사체수거', '포획', '생포', '사살', '인계', '수거']):
        return 'Grade 1'
    if any(k in sp for k in ['수리부엉이', '올빼미', '쇠부엉이', '솔부엉이', '쏙독새']):
        return 'Grade 2'
        
    is_night = (hour >= 20 or hour < 6)
    if is_night:
        if any(k in sp for k in ['종다리', '제비', '황조롱이', '까치', '참새', '멧비둘기']):
            return 'Grade 5'
        if any(k in sp for k in ['오리', '기러기']):
            return 'Grade 4'
        return 'Grade 5'
        
    if any(k in sp for k in ['큰기러기', '쇠기러기', '기러기', '갈매기', '오리', '왜가리', '백로', '황조롱이']):
        return 'Grade 3'
        
    return 'Grade 3'

def determine_sun_window(month, hour):
    if month in [11, 12, 1, 2]:
        if hour in [6, 7]: return 'Dawn'
        elif hour in [16, 17]: return 'Dusk'
        elif 8 <= hour <= 15: return 'Day'
        else: return 'Night'
    elif month in [5, 6, 7, 8]:
        if hour in [4, 5]: return 'Dawn'
        elif hour in [18, 19]: return 'Dusk'
        elif 6 <= hour <= 17: return 'Day'
        else: return 'Night'
    else:
        if hour in [5, 6]: return 'Dawn'
        elif hour in [17, 18]: return 'Dusk'
        elif 7 <= hour <= 16: return 'Day'
        else: return 'Night'

def resample_patrol_2021():
    print(">>> Resampling 2021 Patrol Data...")
    wb = openpyxl.load_workbook(FILE_PATROL_2021, data_only=True)
    ws = wb[wb.sheetnames[0]]
    
    records = []
    for r in range(4, ws.max_row + 1):
        dt_val = ws.cell(row=r, column=3).value
        if not dt_val:
            continue
            
        tm_val = ws.cell(row=r, column=4).value
        species = str(ws.cell(row=r, column=5).value or '미상').strip()
        cnt_val = ws.cell(row=r, column=6).value or 0
        activity = str(ws.cell(row=r, column=7).value or '').strip()
        equip = str(ws.cell(row=r, column=8).value or '').strip()
        inspector = str(ws.cell(row=r, column=9).value or '').strip()
        shooter = str(ws.cell(row=r, column=10).value or '').strip()
        
        shot_7 = int(ws.cell(row=r, column=11).value or 0)
        shot_4 = int(ws.cell(row=r, column=12).value or 0)
        blank = int(ws.cell(row=r, column=13).value or 0)
        
        cat_major = str(ws.cell(row=r, column=14).value or '').strip()
        cat_mid = str(ws.cell(row=r, column=15).value or '').strip()
        area = str(ws.cell(row=r, column=16).value or '').strip()
        area_detail = str(ws.cell(row=r, column=17).value or '').strip()
        map_grp = str(ws.cell(row=r, column=18).value or '').strip()
        
        try:
            count = int(cnt_val)
        except:
            count = 0
            
        if isinstance(dt_val, datetime):
            date_str = dt_val.strftime("%Y-%m-%d")
            month = dt_val.month
        else:
            date_str = str(dt_val)[:10]
            month = int(date_str.split('-')[1]) if '-' in date_str else 1
            
        if isinstance(tm_val, time):
            time_str = tm_val.strftime("%H:%M")
            hour = tm_val.hour
        elif isinstance(tm_val, datetime):
            time_str = tm_val.strftime("%H:%M")
            hour = tm_val.hour
        else:
            time_str = str(tm_val)[:5] if tm_val else "12:00"
            hour = int(time_str.split(':')[0]) if ':' in time_str and time_str.split(':')[0].isdigit() else 12
            
        if count == 0 and ('서식확인' in activity or '순찰' in activity):
            event_type = 'Search_Recon'
        else:
            event_type = 'Actual_Contact'
            
        zone = classify_zone(map_grp, area, area_detail)
        cred = determine_credibility(species, activity, count, hour, date_str)
        sun_win = determine_sun_window(month, hour)
        
        records.append({
            'Year': 2021,
            'Month': month,
            'Date': date_str,
            'Time': time_str,
            'Hour': hour,
            'Species': species,
            'Count': count,
            'Event_Type': event_type,
            'Activity': activity,
            'Equipment': equip,
            'Inspector': inspector,
            'Shooter': shooter,
            'Rounds_7_1_2': shot_7,
            'Rounds_4': shot_4,
            'Rounds_Blank': blank,
            'Rounds_Total': shot_7 + shot_4 + blank,
            'Category_Major': cat_major,
            'Category_Mid': cat_mid,
            'Area_Raw': area,
            'Area_Detail': area_detail,
            'Map_Group': map_grp,
            'Zone_19': zone,
            'Credibility_Grade': cred,
            'Sun_Window': sun_win
        })
        
    df = pd.DataFrame(records)
    os.makedirs(os.path.dirname(OUT_PATROL_CSV), exist_ok=True)
    df.to_csv(OUT_PATROL_CSV, index=False, encoding='utf-8-sig')
    print(f"Patrol 2021 Done: {len(df)} rows saved to {OUT_PATROL_CSV}")
    return df

def resample_strike_2021():
    print(">>> Resampling 2021 Bird Strikes...")
    wb = openpyxl.load_workbook(FILE_STRIKE_21_23, data_only=True)
    ws = wb['목록']
    
    records = []
    for r in range(4, ws.max_row + 1):
        d_val = ws.cell(row=r, column=3).value
        if not d_val:
            continue
        d_str = str(d_val)[:10]
        if not d_str.startswith('2021'):
            continue
            
        t_val = str(ws.cell(row=r, column=4).value or '').strip()
        species = str(ws.cell(row=r, column=5).value or '미상').strip()
        airline = str(ws.cell(row=r, column=6).value or '').strip()
        flight = str(ws.cell(row=r, column=7).value or '').strip()
        model = str(ws.cell(row=r, column=8).value or '').strip()
        route = str(ws.cell(row=r, column=9).value or '').strip()
        runway = str(ws.cell(row=r, column=10).value or '').strip()
        altitude = ws.cell(row=r, column=11).value
        phase = str(ws.cell(row=r, column=12).value or '').strip()
        part = str(ws.cell(row=r, column=13).value or '').strip()
        damage = str(ws.cell(row=r, column=14).value or 'N').strip()
        
        month = int(d_str.split('-')[1]) if '-' in d_str else 1
        hour = int(t_val.split(':')[0]) if ':' in t_val and t_val.split(':')[0].isdigit() else 12
        zone = classify_zone('', runway, '')
        
        records.append({
            'Year': 2021,
            'Month': month,
            'Date': d_str,
            'Time': t_val,
            'Hour': hour,
            'Species': species,
            'Airline': airline,
            'Flight': flight,
            'Aircraft_Model': model,
            'Route': route,
            'Runway_Raw': runway,
            'Zone_19': zone,
            'Altitude_ft': altitude,
            'Flight_Phase': phase,
            'Impact_Part': part,
            'Damage_YN': damage,
            'Is_Damage': 1 if damage == 'Y' else 0,
            'Is_Identified': 0 if species in ['미상', '', '조류미상'] else 1
        })
        
    df = pd.DataFrame(records)
    df.to_csv(OUT_STRIKE_CSV, index=False, encoding='utf-8-sig')
    print(f"Strikes 2021 Done: {len(df)} rows saved to {OUT_STRIKE_CSV}")
    return df

def resample_carcass_2021():
    print(">>> Resampling 2021 Carcass Data...")
    wb = openpyxl.load_workbook(FILE_CARCASS_21_23, data_only=True)
    ws = wb[wb.sheetnames[0]]
    
    records = []
    for r in range(4, ws.max_row + 1):
        d_val = ws.cell(row=r, column=5).value
        if not d_val:
            continue
            
        if isinstance(d_val, datetime):
            d_str = d_val.strftime("%Y-%m-%d")
        else:
            d_str = str(d_val)[:10]
            
        if not d_str.startswith('2021'):
            continue
            
        t_val = ws.cell(row=r, column=6).value
        loc = str(ws.cell(row=r, column=7).value or '').strip()
        species = str(ws.cell(row=r, column=8).value or '미상').strip()
        weather = str(ws.cell(row=r, column=9).value or '').strip()
        status = str(ws.cell(row=r, column=10).value or '').strip()
        remarks = str(ws.cell(row=r, column=11).value or '').strip()
        
        month = int(d_str.split('-')[1]) if '-' in d_str else 1
        
        if isinstance(t_val, time):
            time_str = t_val.strftime("%H:%M")
            hour = t_val.hour
        elif isinstance(t_val, datetime):
            time_str = t_val.strftime("%H:%M")
            hour = t_val.hour
        else:
            time_str = str(t_val)[:5] if t_val else "12:00"
            hour = int(time_str.split(':')[0]) if ':' in time_str and time_str.split(':')[0].isdigit() else 12
            
        is_strike = 1 if any(k in status for k in ['완파', '반파', '충격', '절단']) else 0
        zone = classify_zone('', loc, '')
        
        records.append({
            'Year': 2021,
            'Month': month,
            'Date': d_str,
            'Time': time_str,
            'Hour': hour,
            'Location_Raw': loc,
            'Zone_19': zone,
            'Species': species,
            'Weather': weather,
            'Physical_Condition': status,
            'Is_Strike_Probable': is_strike,
            'Remarks': remarks
        })
        
    df = pd.DataFrame(records)
    df.to_csv(OUT_CARCASS_CSV, index=False, encoding='utf-8-sig')
    print(f"Carcass 2021 Done: {len(df)} rows saved to {OUT_CARCASS_CSV}")
    return df

def main():
    p_df = resample_patrol_2021()
    s_df = resample_strike_2021()
    c_df = resample_carcass_2021()
    
    print("\n==========================================")
    print("=== 2021 YATONGDAE RESAMPLING SUMMARY ===")
    print("==========================================")
    print(f"Total Patrol Rows: {len(p_df)}")
    print(f" - Actual Contact: {len(p_df[p_df['Event_Type'] == 'Actual_Contact'])} ({len(p_df[p_df['Event_Type'] == 'Actual_Contact'])/len(p_df)*100:.1f}%)")
    print(f" - Search Recon: {len(p_df[p_df['Event_Type'] == 'Search_Recon'])} ({len(p_df[p_df['Event_Type'] == 'Search_Recon'])/len(p_df)*100:.1f}%)")
    print(f"Total Rounds Fired: {p_df['Rounds_Total'].sum():,} rounds")
    print(f" - 7 1/2 Rounds: {p_df['Rounds_7_1_2'].sum():,}")
    print(f" - 4 Rounds: {p_df['Rounds_4'].sum():,}")
    print(f" - Blank: {p_df['Rounds_Blank'].sum():,}")
    print(f"Geese Count: {p_df[p_df['Species'].str.contains('기러기')]['Count'].sum():,} birds")
    print("\nTop 5 Zones (Actual Contact):")
    ac_df = p_df[p_df['Event_Type'] == 'Actual_Contact']
    print(ac_df['Zone_19'].value_counts().head(5))
    print(f"\nTotal Bird Strikes (2021): {len(s_df)} (Damaged: {len(s_df[s_df['Is_Damage']==1])})")
    print(f"Total Carcasses (2021): {len(c_df)} (Strike Probable: {len(c_df[c_df['Is_Strike_Probable']==1])})")

if __name__ == '__main__':
    main()
