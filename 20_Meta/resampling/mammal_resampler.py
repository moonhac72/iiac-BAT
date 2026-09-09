import os
import re
import sys
import io
import pandas as pd
import openpyxl
from datetime import datetime, time

# Set stdout encoding
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

FILE_2024 = r"C:\Users\moonh\Desktop\Anythingllm 자료폴더\2024년 자료\2024년 생동물 발생보고.xlsx"
FILE_2025 = r"C:\Users\moonh\Desktop\Anythingllm 자료폴더\2025년 자료\2025년 통계\2025년 생동물.xlsx"
OUT_CSV = r"c:\Users\moonh\Documents\iiac-BAT\_data_raw\resampled\2024_2025_mammal_matrix.csv"

def classify_mammal_zone(loc_str, desc_str=""):
    combined = f"{str(loc_str)} {str(desc_str)}".lower()
    
    if any(k in combined for k in ['aact', '화물터미널', '코압섹', '화물창고', '화물 터미널']):
        return 'ZONE-P-AACT'
    elif any(k in combined for k in ['g2', 'g4', '초소', '외곽', '철책', '울타리']):
        return 'ZONE-P-PERIMETER'
    elif any(k in combined for k in ['841', '842', '851', '852', '800번']):
        return 'ZONE-R4-CARGO'
    elif any(k in combined for k in ['612', '641', '648', '673', '681', '600번', '화물계류장']):
        return 'ZONE-R2-APRON'
    elif any(k in combined for k in ['t2', '253', '200번']):
        return 'ZONE-T2-APRON'
    elif any(k in combined for k in ['t1', '11번', '28번', '51번', '여객']):
        return 'ZONE-T1-APRON'
    elif any(k in combined for k in ['4rw', '4활', '제4활주로', 'ptw']):
        return 'ZONE-R4-AIRSIDE'
    elif any(k in combined for k in ['3rw', '3활', '제3활주로']):
        return 'ZONE-R3-AIRSIDE'
    elif any(k in combined for k in ['1rw', '1활', '제1활주로']):
        return 'ZONE-R1-AIRSIDE'
    elif any(k in combined for k in ['2rw', '2활', '제2활주로']):
        return 'ZONE-R2-AIRSIDE'
    else:
        return 'ZONE-AIRSIDE-GENERAL'

def determine_breach_type(loc_str, desc_str, species):
    combined = f"{str(loc_str)} {str(desc_str)}".lower()
    
    if '창고' in combined or 'aact' in combined or '코압섹' in combined:
        return '화물창고_하역도어_틈새'
    elif '초소' in combined or '철책' in combined or '울타리' in combined or 'g4' in combined or 'g2' in combined:
        return '외곽울타리_하부지면유격'
    elif '차' in combined or '차량' in combined or 'gse' in combined or '조업' in combined:
        return 'GSE조업차량_하부은신'
    elif '기내' in combined or '밀라노' in combined or '화물칸' in combined or 'uld' in combined:
        return '항공기화물_컨테이너유입'
    elif '초지' in combined or '녹지' in combined or '배수로' in combined:
        return '배수로_녹지축_이동경로'
    else:
        return '기동지역_지상개구부'

def parse_2024(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    records = []
    
    # Row 3 is header: [None, '발생일시', '항공사', '이벤트', '생동물구분', '종명', '장소설명', '조치내용', '조치결과']
    for r in range(4, ws.max_row + 1):
        dt_val = ws.cell(row=r, column=2).value
        if not dt_val:
            continue
        
        event = str(ws.cell(row=r, column=4).value or '').strip()
        cat = str(ws.cell(row=r, column=5).value or '').strip()
        species = str(ws.cell(row=r, column=6).value or '').strip()
        loc = str(ws.cell(row=r, column=7).value or '').strip()
        action = str(ws.cell(row=r, column=8).value or '').strip()
        result = str(ws.cell(row=r, column=9).value or '').strip()
        
        # Datetime parse
        if isinstance(dt_val, datetime):
            date_str = dt_val.strftime("%Y-%m-%d")
            time_str = dt_val.strftime("%H:%M")
            hour = dt_val.hour
            month = dt_val.month
        else:
            date_str = str(dt_val)[:10]
            time_str = "00:00"
            hour = 12
            month = int(date_str.split('-')[1]) if '-' in date_str else 1
            
        if not species and cat:
            species = cat
        if not species:
            species = '기타'
            
        zone = classify_mammal_zone(loc, action)
        breach = determine_breach_type(loc, action, species)
        
        # Alert level
        if '고라니' in species:
            alert = '심각'
        elif '사살' in action or '탈출' in action:
            alert = '경계'
        else:
            alert = '주의'
            
        records.append({
            'Year': 2024,
            'Month': month,
            'Date': date_str,
            'Time': time_str,
            'Hour': hour,
            'Species': species,
            'Event_Type': event,
            'Location_Raw': loc,
            'Zone_19': zone,
            'Breach_Type': breach,
            'Action_Detail': action,
            'Result_Detail': result,
            'Alert_Level': alert,
            'Capture_Success': 1 if ('포획' in action or '생포' in action or '사살' in action or '수거' in action or '인수' in action) else 0
        })
    return records

def parse_2025(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    records = []
    
    # Row 3 is header: [None, None, '발생일시', '항공사', '이벤트', '생동물구분', '종명', '장소설명', '조치내용', '조치결과']
    for r in range(4, ws.max_row + 1):
        dt_val = ws.cell(row=r, column=3).value
        if not dt_val:
            continue
            
        event = str(ws.cell(row=r, column=5).value or '').strip()
        cat = str(ws.cell(row=r, column=6).value or '').strip()
        species = str(ws.cell(row=r, column=7).value or '').strip()
        loc = str(ws.cell(row=r, column=8).value or '').strip()
        action = str(ws.cell(row=r, column=9).value or '').strip()
        result = str(ws.cell(row=r, column=10).value or '').strip()
        
        if isinstance(dt_val, datetime):
            date_str = dt_val.strftime("%Y-%m-%d")
            time_str = dt_val.strftime("%H:%M")
            hour = dt_val.hour
            month = dt_val.month
        else:
            date_str = str(dt_val)[:10]
            time_str = "00:00"
            hour = 12
            month = int(date_str.split('-')[1]) if '-' in date_str else 1
            
        if not species and cat:
            species = cat
        if not species:
            species = '기타'
            
        zone = classify_mammal_zone(loc, action)
        breach = determine_breach_type(loc, action, species)
        
        if '고라니' in species:
            alert = '심각'
        elif '사살' in action or '탈출' in action:
            alert = '경계'
        else:
            alert = '주의'
            
        records.append({
            'Year': 2025,
            'Month': month,
            'Date': date_str,
            'Time': time_str,
            'Hour': hour,
            'Species': species,
            'Event_Type': event,
            'Location_Raw': loc,
            'Zone_19': zone,
            'Breach_Type': breach,
            'Action_Detail': action,
            'Result_Detail': result,
            'Alert_Level': alert,
            'Capture_Success': 1 if ('포획' in action or '생포' in action or '사살' in action or '수거' in action or '인수' in action) else 0
        })
    return records

def main():
    rec_2024 = parse_2024(FILE_2024)
    rec_2025 = parse_2025(FILE_2025)
    all_recs = rec_2024 + rec_2025
    df = pd.DataFrame(all_recs)
    
    os.makedirs(os.path.dirname(OUT_CSV), exist_ok=True)
    df.to_csv(OUT_CSV, index=False, encoding='utf-8-sig')
    
    print(f"=== Mammal Resampling Complete ===")
    print(f"Total 2024 records: {len(rec_2024)}")
    print(f"Total 2025 records: {len(rec_2025)}")
    print(f"Total Combined: {len(df)}")
    print(f"Saved to: {OUT_CSV}\n")
    
    print("--- Species Distribution ---")
    print(df.groupby(['Year', 'Species']).size().unstack(fill_value=0))
    print("\n--- Zone Distribution ---")
    print(df.groupby(['Zone_19', 'Species']).size().unstack(fill_value=0))
    print("\n--- Breach Type Distribution ---")
    print(df['Breach_Type'].value_counts())
    print("\n--- Capture Success Rate ---")
    print(df.groupby('Species')['Capture_Success'].agg(['count', 'sum', 'mean']))

if __name__ == '__main__':
    main()
