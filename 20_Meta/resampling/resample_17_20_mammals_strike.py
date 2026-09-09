import os
import openpyxl
import pandas as pd

folder = r"C:\Users\moonh\Desktop\Anythingllm 자료폴더\조류 통계"
out_dir = r"c:\Users\moonh\Documents\iiac-BAT\_data_raw\resampled"
os.makedirs(out_dir, exist_ok=True)

# 1. Carcass
wb_c = openpyxl.load_workbook(os.path.join(folder, "사체수거(17_20년).xlsx"), data_only=True)
ws_c = wb_c[wb_c.sheetnames[0]]
c_rows = []
for r in range(4, ws_c.max_row + 1):
    d_val = ws_c.cell(row=r, column=5).value
    if not d_val: continue
    d_str = str(d_val)[:10]
    status = str(ws_c.cell(row=r, column=10).value or '').strip()
    is_strike = 1 if any(k in status for k in ['완파', '반파', '충격', '절단']) else 0
    c_rows.append({
        'Date': d_str,
        'Year': int(d_str[:4]),
        'Month': int(d_str[5:7]),
        'Time': str(ws_c.cell(row=r, column=6).value or '')[:8],
        'Location_Raw': str(ws_c.cell(row=r, column=7).value or '').strip(),
        'Species': str(ws_c.cell(row=r, column=8).value or '').strip(),
        'Weather': str(ws_c.cell(row=r, column=9).value or '').strip(),
        'Physical_Condition': status,
        'Is_Strike_Probable': is_strike,
        'Remark': str(ws_c.cell(row=r, column=11).value or '').strip()
    })
df_c = pd.DataFrame(c_rows)
p_c = os.path.join(out_dir, "2017_2020_carcass_resampled.csv")
df_c.to_csv(p_c, index=False, encoding='utf-8-sig')
print(f"Saved Carcass: {len(df_c)} rows -> {p_c}")

# 2. Live Mammals / Animals
wb_l = openpyxl.load_workbook(os.path.join(folder, "생동물(17_20년).xlsx"), data_only=True)
ws_l = wb_l[wb_l.sheetnames[0]]
l_rows = []
for r in range(4, ws_l.max_row + 1):
    d_val = ws_l.cell(row=r, column=3).value
    if not d_val: continue
    d_str = str(d_val)[:10]
    l_rows.append({
        'DateTime': str(d_val)[:19],
        'Date': d_str,
        'Year': int(d_str[:4]),
        'Month': int(d_str[5:7]),
        'Time': str(d_val)[11:16],
        'Species': str(ws_l.cell(row=r, column=6).value or '').strip(),
        'Event': str(ws_l.cell(row=r, column=5).value or '').strip(),
        'Action': str(ws_l.cell(row=r, column=7).value or '').strip(),
        'Location_Raw': str(ws_l.cell(row=r, column=8).value or '').strip(),
        'Action_Detail': str(ws_l.cell(row=r, column=9).value or '').strip(),
        'Result': str(ws_l.cell(row=r, column=10).value or '').strip()
    })
df_l = pd.DataFrame(l_rows)
p_l = os.path.join(out_dir, "2017_2020_live_mammals_resampled.csv")
df_l.to_csv(p_l, index=False, encoding='utf-8-sig')
print(f"Saved Live Animals: {len(df_l)} rows -> {p_l}")

# 3. Bird Strike
wb_s = openpyxl.load_workbook(os.path.join(folder, "조류충돌(17_20년).xlsx"), data_only=True)
ws_s = wb_s['목록']
s_rows = []
for r in range(4, ws_s.max_row + 1):
    d_val = ws_s.cell(row=r, column=3).value
    if not d_val: continue
    d_str = str(d_val)[:10]
    dmg = str(ws_s.cell(row=r, column=14).value or 'N').strip()
    sp = str(ws_s.cell(row=r, column=5).value or '미상').strip()
    s_rows.append({
        'Date': d_str,
        'Year': int(d_str[:4]),
        'Month': int(d_str[5:7]),
        'Time': str(ws_s.cell(row=r, column=4).value or '')[:8],
        'Species': sp,
        'Is_Identified': 0 if sp in ['미상', '', '조류미상'] else 1,
        'Airline': str(ws_s.cell(row=r, column=6).value or '').strip(),
        'Flight': str(ws_s.cell(row=r, column=7).value or '').strip(),
        'Aircraft_Model': str(ws_s.cell(row=r, column=8).value or '').strip(),
        'Route': str(ws_s.cell(row=r, column=9).value or '').strip(),
        'Runway_Raw': str(ws_s.cell(row=r, column=10).value or '').strip(),
        'Altitude_ft': ws_s.cell(row=r, column=11).value,
        'Flight_Phase': str(ws_s.cell(row=r, column=12).value or '').strip(),
        'Impact_Part': str(ws_s.cell(row=r, column=13).value or '').strip(),
        'Damage_YN': dmg,
        'Is_Damage': 1 if dmg == 'Y' else 0,
        'Position': str(ws_s.cell(row=r, column=15).value or '').strip(),
        'Found_Carcass': str(ws_s.cell(row=r, column=16).value or '').strip(),
        'DNA_Requested': str(ws_s.cell(row=r, column=17).value or '').strip()
    })
df_s = pd.DataFrame(s_rows)
p_s = os.path.join(out_dir, "2017_2020_birdstrike_resampled.csv")
df_s.to_csv(p_s, index=False, encoding='utf-8-sig')
print(f"Saved Strikes: {len(df_s)} rows -> {p_s}")
