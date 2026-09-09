import os
import re
import sys
import io
import pandas as pd
import openpyxl
from datetime import datetime, time

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

FILE_CARCASS_2024 = r"C:\Users\moonh\Desktop\Anythingllm 자료폴더\2024년 자료\2024년 사체수거 현황.xlsx"
FILE_CARCASS_2025 = r"C:\Users\moonh\Desktop\Anythingllm 자료폴더\2025년 자료\2025년 통계\2025년 기동지역 사체수거.xlsx"
OUT_CSV = r"c:\Users\moonh\Documents\iiac-BAT\_data_raw\resampled\2024_2025_carcass_matrix.csv"

def classify_carcass_zone(loc_str):
    s = str(loc_str).lower()
    if '34l' in s or '34' in s and '남' in s:
        return 'ZONE-R4-S'
    elif '16r' in s or '16' in s and '북' in s and ('4rw' in s or 'p' in s):
        return 'ZONE-R4-N'
    elif '34r' in s or ('3rw' in s and ('w' in s or 'u' in s or '남' in s)):
        return 'ZONE-R3-S'
    elif '16l' in s or ('3rw' in s and ('n' in s or '북' in s)):
        return 'ZONE-R3-N'
    elif '33l' in s or ('2rw' in s and ('d' in s or '남' in s)):
        return 'ZONE-R2-S'
    elif '15r' in s or ('2rw' in s and ('b' in s or '북' in s)):
        return 'ZONE-R2-N'
    elif '33r' in s or ('1rw' in s and ('c' in s or 'e' in s or 'k' in s or '남' in s)):
        return 'ZONE-R1-S'
    elif '15l' in s or ('1rw' in s and ('a' in s or '북' in s)):
        return 'ZONE-R1-N'
    elif '4rw' in s or '4활' in s:
        return 'ZONE-R4-AIRSIDE'
    elif '3rw' in s or '3활' in s:
        return 'ZONE-R3-AIRSIDE'
    elif '2rw' in s or '2활' in s:
        return 'ZONE-R2-AIRSIDE'
    elif '1rw' in s or '1활' in s:
        return 'ZONE-R1-AIRSIDE'
    elif any(k in s for k in ['600', '612', '641', '648', '673', '841', '851', 'aact']):
        return 'ZONE-R2-APRON'
    elif any(k in s for k in ['200', '217', '251', '253', 't2']):
        return 'ZONE-T2-APRON'
    elif any(k in s for k in ['t1', '탑승동', '11번', '28번']):
        return 'ZONE-T1-APRON'
    else:
        return 'ZONE-AIRSIDE-GENERAL'

def parse_carcass_2024(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    records = []
    # Row 3 is header: [None, None, '일자', '시간', '수거위치', '조류명', '날씨', '상태', '비고']
    for r in range(4, ws.max_row + 1):
        dt_val = ws.cell(row=r, column=3).value
        if not dt_val:
            continue
        tm_val = ws.cell(row=r, column=4).value
        loc = str(ws.cell(row=r, column=5).value or '').strip()
        species = str(ws.cell(row=r, column=6).value or '').strip()
        weather = str(ws.cell(row=r, column=7).value or '').strip()
        status = str(ws.cell(row=r, column=8).value or '').strip()
        remarks = str(ws.cell(row=r, column=9).value or '').strip()
        
        if isinstance(dt_val, datetime):
            date_str = dt_val.strftime("%Y-%m-%d")
            month = dt_val.month
        else:
            date_str = str(dt_val)[:10]
            month = int(date_str.split('-')[1]) if '-' in date_str else 1
            
        if isinstance(tm_val, time):
            time_str = tm_val.strftime("%H:%M")
            hour = tm_val.hour
        elif isinstance(tm_val, datetime):
            time_str = tm_val.strftime("%H:%M")
            hour = tm_val.hour
        else:
            time_str = str(tm_val)[:5] if tm_val else "12:00"
            hour = int(time_str.split(':')[0]) if ':' in time_str and time_str.split(':')[0].isdigit() else 12
            
        # Parse physical condition
        is_strike_probable = 1 if any(k in status for k in ['완파', '반파', '파손', '충격', '절단']) else 0
        
        records.append({
            'Year': 2024,
            'Month': month,
            'Date': date_str,
            'Time': time_str,
            'Hour': hour,
            'Location_Raw': loc,
            'Zone_19': classify_carcass_zone(loc),
            'Species': species if species else '미상',
            'Weather': weather,
            'Physical_Condition': status,
            'Is_Strike_Probable': is_strike_probable,
            'Remarks': remarks
        })
    return records

def parse_carcass_2025(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    records = []
    # Row 3 is header: [None, None, None, '일자', '시간', '수거위치', '조류명', '날씨', '상태', '비고']
    for r in range(4, ws.max_row + 1):
        dt_val = ws.cell(row=r, column=4).value
        if not dt_val:
            continue
        tm_val = ws.cell(row=r, column=5).value
        loc = str(ws.cell(row=r, column=6).value or '').strip()
        species = str(ws.cell(row=r, column=7).value or '').strip()
        weather = str(ws.cell(row=r, column=8).value or '').strip()
        status = str(ws.cell(row=r, column=9).value or '').strip()
        remarks = str(ws.cell(row=r, column=10).value or '').strip()
        
        if isinstance(dt_val, datetime):
            date_str = dt_val.strftime("%Y-%m-%d")
            month = dt_val.month
        else:
            date_str = str(dt_val)[:10]
            month = int(date_str.split('-')[1]) if '-' in date_str else 1
            
        if isinstance(tm_val, time):
            time_str = tm_val.strftime("%H:%M")
            hour = tm_val.hour
        elif isinstance(tm_val, datetime):
            time_str = tm_val.strftime("%H:%M")
            hour = tm_val.hour
        else:
            time_str = str(tm_val)[:5] if tm_val else "12:00"
            hour = int(time_str.split(':')[0]) if ':' in time_str and time_str.split(':')[0].isdigit() else 12
            
        is_strike_probable = 1 if any(k in status for k in ['완파', '반파', '파손', '충격', '절단']) else 0
        
        records.append({
            'Year': 2025,
            'Month': month,
            'Date': date_str,
            'Time': time_str,
            'Hour': hour,
            'Location_Raw': loc,
            'Zone_19': classify_carcass_zone(loc),
            'Species': species if species else '미상',
            'Weather': weather,
            'Physical_Condition': status,
            'Is_Strike_Probable': is_strike_probable,
            'Remarks': remarks
        })
    return records

def main():
    rec_2024 = parse_carcass_2024(FILE_CARCASS_2024)
    rec_2025 = parse_carcass_2025(FILE_CARCASS_2025)
    all_recs = rec_2024 + rec_2025
    df = pd.DataFrame(all_recs)
    
    os.makedirs(os.path.dirname(OUT_CSV), exist_ok=True)
    df.to_csv(OUT_CSV, index=False, encoding='utf-8-sig')
    
    print(f"=== Carcass Resampling Complete ===")
    print(f"Total 2024 records: {len(rec_2024)}")
    print(f"Total 2025 records: {len(rec_2025)}")
    print(f"Total Combined: {len(df)}")
    print(f"Saved to: {OUT_CSV}\n")
    
    print("--- Top 10 Species Recovered ---")
    print(df['Species'].value_counts().head(10))
    print("\n--- Physical Condition ---")
    print(df['Physical_Condition'].value_counts().head(10))
    print("\n--- Strike Probable vs Non-Strike ---")
    print(df.groupby('Year')['Is_Strike_Probable'].value_counts().unstack(fill_value=0))
    print("\n--- Top Zones for Carcass Recovery ---")
    print(df['Zone_19'].value_counts())

if __name__ == '__main__':
    main()
